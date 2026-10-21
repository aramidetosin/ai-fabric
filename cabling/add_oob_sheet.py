import sys
#!/usr/bin/env python3
"""Add the OOB sheet to the operator's uploaded workbook and refresh the Summary. Their Backend/Frontend
sheets and their Summary wording are left exactly as they are; only the OOB row, the TOTAL and the
out-of-scope line change."""
import csv
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from gen_oob import oob_rows

SRC = sys.argv[1] if len(sys.argv) > 1 else "ai-fabric-cabling.xlsx"
OUT = "ai-fabric-cabling.xlsx"
HDR_BG, HDR_FG = "1F3864", "FFFFFF"
BAND = "DEEBF7"                       # OOB is a single speed, so one calm band rather than a speed legend
thin = Side(style="thin", color="BFBFBF")

wb = load_workbook(SRC)
if "OOB" in wb.sheetnames: del wb["OOB"]

hdr = ["a_device", "a_port", "b_device", "b_port", "speed", "media", "notes"]
rows = oob_rows()
ws = wb.create_sheet("OOB")
ws.append(["OOB / BMC: one BMC port per server, one management port per backend switch, two per FlashBlade appliance. "
           "<OOB subnet>/24, gateway on the firewall. oob-sw model and port count still to be selected; ports are sequential."])
ws["A1"].font = Font(italic=True, color="595959")
ws.append([])
ws.append(hdr)
for r in rows: ws.append(list(r))
for c in range(1, len(hdr) + 1):
    cell = ws.cell(row=3, column=c)
    cell.font = Font(bold=True, color=HDR_FG); cell.fill = PatternFill("solid", fgColor=HDR_BG)
    cell.alignment = Alignment(vertical="center")
for r in range(4, 4 + len(rows)):
    for c in range(1, len(hdr) + 1):
        cell = ws.cell(row=r, column=c)
        cell.border = Border(bottom=thin); cell.fill = PatternFill("solid", fgColor=BAND)
ws.freeze_panes = ws.cell(row=4, column=1)
ws.auto_filter.ref = f"A3:{get_column_letter(len(hdr))}{3 + len(rows)}"
for c in range(1, len(hdr) + 1):
    width = max(len(str(ws.cell(row=r, column=c).value or "")) for r in range(3, 4 + len(rows)))
    ws.column_dimensions[get_column_letter(c)].width = min(width + 3, 62)

# ---- Summary: insert an OOB line above TOTAL, refresh TOTAL, amend the scope note
s = wb["Summary"]
s.insert_rows(5)
n_sw = sum(1 for r in rows if r[0].startswith("be-"))
n_gpu = sum(1 for r in rows if r[0].startswith("gpu"))
n_fb = sum(1 for r in rows if r[0].startswith("flashblade"))
n_cpu = len(rows) - n_sw - n_gpu - n_fb
s.cell(row=5, column=1, value="OOB / BMC")
s.cell(row=5, column=2, value=len(rows))
s.cell(row=5, column=3, value=f"{n_sw} backend switches + {n_gpu} GPU nodes + {n_cpu} CPU nodes (ctrl1-6 + bastion) "
                             f"+ {n_fb} FlashBlade mgmt, all 1G RJ45 to oob-sw")
total = 384 + 86 + len(rows)
s.cell(row=6, column=2, value=total)
s.cell(row=6, column=3, value="every GPU node: 8 x 400G backend rails + 2 x 100G frontend + 1 x 1G BMC = 11 links")
for c in (1, 2): s.cell(row=6, column=c).font = Font(bold=True)
s.cell(row=8, column=1, value="Out of scope: border leaves, firewall, frontend leaf-spine uplinks (ports 31/32), "
                              "frontend switch management, oob-sw uplink to the firewall.")
wb.save(OUT)
print(f"  OOB sheet added: {len(rows)} links | workbook total now {total}")
